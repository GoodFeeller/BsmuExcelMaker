from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import openpyxl

from app.cell_utils import normalize_cell
from app.modules.duty_summary.models import CountSummaryStats

DATA_START_ROW = 2


@dataclass(frozen=True)
class CountColumns:
    offset: int = 0

    @property
    def room(self) -> int:
        return 1 + self.offset

    @property
    def date(self) -> int:
        return 2 + self.offset

    @property
    def students(self) -> int:
        return 4 + self.offset

    @property
    def program(self) -> int:
        return 5 + self.offset

    @property
    def plan(self) -> int:
        return 8 + self.offset

    @property
    def fact(self) -> int:
        return 10 + self.offset

    @property
    def cushion(self) -> int:
        return 11 + self.offset

    @property
    def categories(self) -> tuple[int, ...]:
        return tuple(12 + index + self.offset for index in range(8))

    @property
    def other(self) -> int:
        return 19 + self.offset

    @property
    def max_col(self) -> int:
        return self.other


def _is_numeric(value) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _format_number(value) -> str:
    if value is None:
        return "0"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(value).replace(".", ",")
    return str(value)


def _header_text(worksheet, row: int, col: int) -> str:
    return normalize_cell(worksheet.cell(row, col).value).casefold()


def _detect_columns(worksheet) -> CountColumns:
    for header_row in (1, 2):
        for col in (1, 2):
            if _header_text(worksheet, header_row, col).startswith("кабинет"):
                return CountColumns(offset=col - 1)

    if _header_text(worksheet, 1, 2).startswith("дата") and worksheet.cell(
        DATA_START_ROW, 1
    ).value not in (None, ""):
        return CountColumns(offset=0)

    raise ValueError(
        "Ожидается файл «2. Подсчет.xlsx»: не найден столбец «Кабинет» в заголовке"
    )


def _validate_headers(worksheet, columns: CountColumns) -> None:
    header_row = 1
    room_ok = _header_text(worksheet, header_row, columns.room).startswith("кабинет")
    date_ok = _header_text(worksheet, header_row, columns.date).startswith("дата")
    if not room_ok and not date_ok:
        raise ValueError(
            "Ожидается файл «2. Подсчет.xlsx»: не найдены столбцы «Кабинет» или «Дата»"
        )
    if "обуча" not in _header_text(worksheet, header_row, columns.students):
        raise ValueError(
            "Ожидается файл «2. Подсчет.xlsx»: проверьте заголовок «Обучающиеся»"
        )
    if not _header_text(worksheet, header_row, columns.plan).startswith("план"):
        raise ValueError(
            "Ожидается файл «2. Подсчет.xlsx»: проверьте заголовок «План»"
        )


def _collect_totals(worksheet, columns: CountColumns, data_end: int) -> dict[int, object]:
    total_columns = (
        columns.plan,
        columns.fact,
        columns.cushion,
        *columns.categories,
    )
    totals: dict[int, object] = {col: None for col in total_columns}

    primary_row = None
    for row in range(worksheet.max_row, data_end, -1):
        if worksheet.cell(row, columns.room).value not in (None, ""):
            continue
        fact = worksheet.cell(row, columns.fact).value
        cushion = worksheet.cell(row, columns.cushion).value
        if _is_numeric(fact) and _is_numeric(cushion):
            primary_row = row
            break

    if primary_row is not None:
        for col in total_columns:
            value = worksheet.cell(primary_row, col).value
            if _is_numeric(value):
                totals[col] = value

    for row in range(worksheet.max_row, data_end, -1):
        if worksheet.cell(row, columns.room).value not in (None, ""):
            continue
        for col in total_columns:
            if totals[col] is not None:
                continue
            value = worksheet.cell(row, col).value
            if _is_numeric(value):
                totals[col] = value

    if totals[columns.fact] is None:
        raise ValueError("Не найдена строка итогов в файле подсчёта")
    return totals


def _find_data_end_from_totals(worksheet, columns: CountColumns) -> int:
    for row in range(worksheet.max_row, DATA_START_ROW - 1, -1):
        if worksheet.cell(row, columns.room).value not in (None, ""):
            return row
    return DATA_START_ROW - 1


def _count_groups(worksheet, data_end: int, columns: CountColumns) -> int:
    seen_int: set[tuple[str, int | float]] = set()
    seen_c4_only: set[str] = set()
    count = 0
    for row in range(DATA_START_ROW, data_end + 1):
        col4 = worksheet.cell(row, columns.students).value
        col5 = worksheet.cell(row, columns.program).value
        if not col4:
            continue
        label = str(col4).strip()
        if _is_numeric(col5):
            key = (label, col5)
            if key not in seen_int:
                seen_int.add(key)
                count += 1
        else:
            if label not in seen_c4_only:
                seen_c4_only.add(label)
                count += 1
    return count


def _count_visits(worksheet, data_end: int, columns: CountColumns) -> int:
    return sum(
        1
        for row in range(DATA_START_ROW, data_end + 1)
        if worksheet.cell(row, columns.room).value not in (None, "")
    )


def _format_prochee(
    worksheet,
    data_end: int,
    columns: CountColumns,
    sum_value,
) -> str | int | float | None:
    lines: list[str] = []
    for row in range(DATA_START_ROW, data_end + 1):
        value = worksheet.cell(row, columns.other).value
        if value in (None, ""):
            continue
        col4 = worksheet.cell(row, columns.students).value
        label = str(col4).strip() if col4 else ""
        if _is_numeric(value) and value > 0 and len(label) > 2:
            lines.append(f"{_format_number(value)}-{label}")
        elif not _is_numeric(value):
            lines.append(str(value).strip())

    if lines:
        return "\n".join(lines)
    return sum_value


def parse_count_workbook(path: str | Path) -> CountSummaryStats:
    workbook = openpyxl.load_workbook(path, data_only=True)
    try:
        worksheet = workbook.active
        columns = _detect_columns(worksheet)
        _validate_headers(worksheet, columns)

        data_end = _find_data_end_from_totals(worksheet, columns)
        if data_end < DATA_START_ROW:
            raise ValueError("В файле подсчёта нет строк с данными")

        totals = _collect_totals(worksheet, columns, data_end)

        report_date = ""
        for row in range(DATA_START_ROW, data_end + 1):
            raw = worksheet.cell(row, columns.date).value
            if raw is not None and normalize_cell(raw):
                if hasattr(raw, "strftime"):
                    report_date = raw.strftime("%d.%m.%Y")
                else:
                    report_date = normalize_cell(raw)
                break

        category_cols = columns.categories
        other_value = _format_prochee(
            worksheet,
            data_end,
            columns,
            totals[columns.other],
        )

        return CountSummaryStats(
            report_date=report_date,
            groups_count=_count_groups(worksheet, data_end, columns),
            visits_count=_count_visits(worksheet, data_end, columns),
            plan_total=totals[columns.plan],
            fact_total=totals[columns.fact],
            cushion_total=totals[columns.cushion],
            ko=totals[category_cols[0]],
            i_cat=totals[category_cols[1]],
            pk=totals[category_cols[2]],
            pp=totals[category_cols[3]],
            attestation=totals[category_cols[4]],
            bgmu=totals[category_cols[5]],
            mfiu=totals[category_cols[6]],
            other=other_value,
            source_path=str(path),
        )
    finally:
        workbook.close()
