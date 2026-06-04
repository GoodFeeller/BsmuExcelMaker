from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt

from app.app_paths import app_root
from app.modules.duty_summary.models import CountSummaryStats

TABLE_FONT_NAME = "Times New Roman"
TABLE_FONT_SIZE_PT = 12
REPORT_HEADER_FONT_SIZE_PT = 14


def resolve_summary_template_path() -> Path:
    bundled = app_root() / "assets" / "summary_template.docx"
    if bundled.exists():
        return bundled
    raise FileNotFoundError(
        "Не найден шаблон сводной информации. Положите файл в assets/summary_template.docx"
    )


def default_summary_output_path(count_path: str | Path) -> Path:
    directory = Path(count_path).resolve().parent
    return directory / "3. Сводная информация.docx"


def _format_cell_value(value) -> str:
    if value is None:
        return "0"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(value).replace(".", ",")
    return str(value).strip() or "0"


def _apply_cell_font(cell, size_pt: int = TABLE_FONT_SIZE_PT) -> None:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.name = TABLE_FONT_NAME
            run.font.size = Pt(size_pt)


def _set_cell_text(cell, text: str) -> None:
    cell.text = text or ""
    _apply_cell_font(cell)


def _set_paragraph_text(paragraph, text: str, *, size_pt: int = REPORT_HEADER_FONT_SIZE_PT) -> None:
    paragraph.text = text
    for run in paragraph.runs:
        run.font.name = TABLE_FONT_NAME
        run.font.size = Pt(size_pt)


def _replace_paragraph_starting_with(
    document: Document,
    prefix: str,
    new_text: str,
    *,
    size_pt: int = REPORT_HEADER_FONT_SIZE_PT,
) -> None:
    prefix_fold = prefix.casefold()
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text.casefold().startswith(prefix_fold):
            _set_paragraph_text(paragraph, new_text, size_pt=size_pt)
            return


def _replace_students_line(document: Document, cushion_text: str) -> None:
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        folded = text.casefold()
        if folded.startswith("из ") and "обучающихся" in folded:
            _set_paragraph_text(
                paragraph,
                f"Из {cushion_text} обучающихся:",
            )
            return


def _replace_duty_specialist_line(document: Document, specialist: str) -> None:
    if not specialist.strip():
        return

    paragraphs = document.paragraphs
    for index, paragraph in enumerate(paragraphs):
        folded = paragraph.text.casefold()
        if "специалист по обеспечению" not in folded:
            continue

        _set_paragraph_text(paragraph, "Специалист по обеспечению")
        if index + 1 < len(paragraphs):
            next_paragraph = paragraphs[index + 1]
            next_folded = next_paragraph.text.casefold()
            if "симуляционн" in next_folded or "\t" in next_paragraph.text:
                label = next_paragraph.text.split("\t", 1)[0].strip()
                label = label or "симуляционного обучения"
                _set_paragraph_text(
                    next_paragraph,
                    f"{label}\t{specialist}",
                )
                return

        _set_paragraph_text(
            paragraph,
            f"Специалист по обеспечению\t\t{specialist}",
        )
        return

    for paragraph in paragraphs:
        folded = paragraph.text.casefold()
        if "ответственный специалист" in folded:
            label = (
                paragraph.text.split("\t", 1)[0].strip()
                if "\t" in paragraph.text
                else "Ответственный специалист"
            )
            _set_paragraph_text(paragraph, f"{label}\t{specialist}")
            return


def export_summary_to_word(
    stats: CountSummaryStats,
    *,
    report_date: str,
    duty_specialist: str,
    output_path: str | Path,
) -> None:
    template_path = resolve_summary_template_path()
    document = Document(str(template_path))

    _replace_paragraph_starting_with(
        document,
        "по состоянию на",
        f"по состоянию на {report_date}",
    )

    cushion_text = _format_cell_value(stats.cushion_total)
    _replace_students_line(document, cushion_text)

    _replace_duty_specialist_line(document, duty_specialist)

    if len(document.tables) < 2:
        raise ValueError("В шаблоне сводной информации ожидаются две таблицы")

    overview = document.tables[0]
    if len(overview.rows) < 2:
        raise ValueError("В шаблоне не найдена строка сводной таблицы")
    overview_row = overview.rows[1].cells
    values = [
        report_date,
        _format_cell_value(stats.groups_count),
        _format_cell_value(stats.visits_count),
        _format_cell_value(stats.plan_total),
        _format_cell_value(stats.fact_total),
        _format_cell_value(stats.cushion_total),
    ]
    for cell, value in zip(overview_row, values):
        _set_cell_text(cell, value)

    categories = document.tables[1]
    if len(categories.rows) < 3:
        raise ValueError("В шаблоне не найдена строка категорий")
    category_cells = categories.rows[2].cells
    for cell, value in zip(category_cells, stats.category_values):
        _set_cell_text(cell, _format_cell_value(value))

    document.save(str(output_path))
