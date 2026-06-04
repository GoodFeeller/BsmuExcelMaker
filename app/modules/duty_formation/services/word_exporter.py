from datetime import date
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table

from app.app_paths import app_root
from app.cell_utils import normalize_cell


def _teacher_key(value: str) -> str:
    return normalize_cell(value).strip().casefold()


def _teacher_merge_spans(teachers: list[str]) -> list[tuple[int, int]]:
    spans: list[tuple[int, int]] = []
    index = 0
    total = len(teachers)

    while index < total:
        anchor = _teacher_key(teachers[index])
        if not anchor:
            index += 1
            continue

        end = index + 1
        while end < total:
            current = _teacher_key(teachers[end])
            if current and current != anchor:
                break
            end += 1

        if end - index > 1:
            spans.append((index, end))
        index = end

    return spans
from app.modules.duty_formation.models import GroupAggregate, ScheduleLine
from app.modules.duty_formation.services.responsible_mapping import (
    responsible_merge_spans_for_group,
)
from app.modules.duty_formation.services.schedule_aggregator import (
    _schedule_blocks,
    _schedule_line_has_content,
)

TABLE_FONT_NAME = "Times New Roman"
TABLE_FONT_SIZE_PT = 12
REPORT_HEADER_FONT_SIZE_PT = 14


def resolve_template_path() -> Path:
    bundled = app_root() / "assets" / "word_template.docx"
    if bundled.exists():
        return bundled

    root = app_root()
    candidates = sorted(
        path
        for path in root.glob("*.docx")
        if not path.name.startswith("~$")
    )
    if candidates:
        return candidates[0]

    raise FileNotFoundError(
        "Не найден шаблон Word. Положите файл в assets/word_template.docx"
    )


def _apply_cell_font(
    cell,
    font_name: str = TABLE_FONT_NAME,
    size_pt: int = TABLE_FONT_SIZE_PT,
) -> None:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.name = font_name
            run.font.size = Pt(size_pt)


def _apply_paragraph_font(
    paragraph,
    font_name: str = TABLE_FONT_NAME,
    size_pt: int = TABLE_FONT_SIZE_PT,
) -> None:
    for run in paragraph.runs:
        run.font.name = font_name
        run.font.size = Pt(size_pt)


def _trim_trailing_empty_lines(lines: list[str]) -> list[str]:
    values = list(lines) if lines else []
    while len(values) > 1 and not values[-1].strip():
        values.pop()
    return values if values else [""]


def _clear_extra_paragraphs(cell, keep_count: int) -> None:
    while len(cell.paragraphs) > keep_count:
        element = cell.paragraphs[-1]._element
        element.getparent().remove(element)


WORD_BORDER_SIZE = 4  # 0.5 pt в единицах Word (1/8 pt)


def _set_cell_bottom_border(
    cell,
    size: int = WORD_BORDER_SIZE,
    color: str = "475569",
) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for existing in borders.findall(qn("w:bottom")):
        borders.remove(existing)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "0")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def _set_cell_text(cell, text: str) -> None:
    cell.text = text or ""
    _apply_cell_font(cell)


def _merge_column_cells(table: Table, col: int, row_start: int, row_end: int) -> None:
    if row_end <= row_start:
        return
    top = table.rows[row_start].cells[col]
    for row_index in range(row_start + 1, row_end + 1):
        top.merge(table.rows[row_index].cells[col])


def _append_group_rows(
    table: Table,
    group: GroupAggregate,
    *,
    group_border_after: bool = False,
) -> None:
    lines = [
        line
        for line in (group.lines or [])
        if _schedule_line_has_content(
            line.time,
            line.teacher,
            line.responsible,
            line.room,
        )
    ]
    if not lines:
        return

    row_start = len(table.rows)

    for line in lines:
        row = table.add_row()
        cells = row.cells
        _set_cell_text(cells[2], line.time)
        _set_cell_text(cells[3], line.teacher)
        _set_cell_text(cells[4], line.responsible)
        _set_cell_text(cells[5], line.room)
        if line.border_after:
            for cell in cells:
                _set_cell_bottom_border(cell)

    if group_border_after:
        for cell in table.rows[-1].cells:
            _set_cell_bottom_border(cell)

    row_end = len(table.rows) - 1
    first_row = table.rows[row_start]
    _set_cell_text(first_row.cells[0], str(group.index))
    _set_cell_text(first_row.cells[1], group.group)
    _merge_column_cells(table, 0, row_start, row_end)
    _merge_column_cells(table, 1, row_start, row_end)

    times = [line.time for line in lines]
    teachers = [line.teacher for line in lines]
    responsibles = [line.responsible for line in lines]
    for block_start, block_end in _schedule_blocks(times):
        if block_end - block_start <= 1:
            continue
        block_row_start = row_start + block_start
        block_row_end = row_start + block_end - 1
        _merge_column_cells(table, 2, block_row_start, block_row_end)

    for rel_start, rel_end in _teacher_merge_spans(teachers):
        if rel_end - rel_start <= 1:
            continue
        _merge_column_cells(
            table,
            3,
            row_start + rel_start,
            row_start + rel_end - 1,
        )

    for rel_start, rel_end in responsible_merge_spans_for_group(responsibles):
        if rel_end - rel_start <= 1:
            continue
        _merge_column_cells(
            table,
            4,
            row_start + rel_start,
            row_start + rel_end - 1,
        )


def _fill_cell(cell, lines: list[str]) -> None:
    values = _trim_trailing_empty_lines(lines)
    cell.text = values[0]
    for line in values[1:]:
        cell.add_paragraph(line)
    _clear_extra_paragraphs(cell, len(values))
    _apply_cell_font(cell)


def _remove_table_rows(table: Table, keep_rows: int) -> None:
    while len(table.rows) > keep_rows:
        row_element = table.rows[-1]._tr
        table._tbl.remove(row_element)


def _apply_table_font(table: Table) -> None:
    for row in table.rows:
        for cell in row.cells:
            _apply_cell_font(cell)


def _normalize_report_phrases(document: Document) -> None:
    for paragraph in document.paragraphs:
        text = paragraph.text
        if "РЦПА и СО" in text or "РЦПА И СО" in text:
            paragraph.text = (
                text.replace("РЦПА и СО", "РЦПАиСО").replace("РЦПА и СО", "РЦПАиСО")
            )


def _update_status_date(document: Document, report_date: str) -> None:
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("По состоянию на") or text.startswith("по состоянию на"):
            paragraph.text = f"по состоянию на {report_date}"
            _apply_paragraph_font(paragraph, size_pt=REPORT_HEADER_FONT_SIZE_PT)
            return


def _update_duty_specialist(document: Document, duty_specialist: str) -> None:
    if not duty_specialist:
        return
    for paragraph in document.paragraphs:
        if paragraph.text.startswith("Специалист по обеспечению"):
            paragraph.text = f"Специалист по обеспечению\t\t{duty_specialist}"
            _apply_paragraph_font(paragraph, size_pt=REPORT_HEADER_FONT_SIZE_PT)
            return


def export_groups_to_word(
    groups: list[GroupAggregate],
    output_path: str | Path,
    report_date: str | None = None,
    duty_specialist: str | None = None,
) -> None:
    if not groups:
        raise ValueError("Нет данных для экспорта")

    template_path = resolve_template_path()
    document = Document(str(template_path))
    if not document.tables:
        raise ValueError("В шаблоне Word не найдена таблица")

    _normalize_report_phrases(document)
    date_text = report_date or date.today().strftime("%d.%m.%Y")
    _update_status_date(document, date_text)
    _update_duty_specialist(document, duty_specialist or "")
    table = document.tables[0]
    _remove_table_rows(table, keep_rows=1)

    last_group_index = len(groups) - 1
    for group_index, group in enumerate(groups):
        _append_group_rows(
            table,
            group,
            group_border_after=group_index < last_group_index,
        )

    _apply_table_font(table)
    document.save(str(output_path))
