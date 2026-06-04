from __future__ import annotations

from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from app.app_paths import app_root
from app.modules.duty_formation.services.count_row_builder import COUNT_DATA_COLUMNS

DATA_START_ROW = 2
TEMPLATE_STYLE_ROW = 2


def resolve_count_template_path() -> Path:
    bundled = app_root() / "assets" / "count_template.xlsx"
    if bundled.exists():
        return bundled
    raise FileNotFoundError(
        "Не найден шаблон подсчёта. Положите файл в assets/count_template.xlsx"
    )


def default_count_output_path(word_path: str | Path) -> Path:
    directory = Path(word_path).resolve().parent
    return directory / "2. Подсчет.xlsx"


def _copy_row_style(ws, source_row: int, target_row: int, max_col: int) -> None:
    for col in range(1, max_col + 1):
        source = ws.cell(source_row, col)
        target = ws.cell(target_row, col)
        if source.has_style:
            target.font = copy(source.font)
            target.border = copy(source.border)
            target.fill = copy(source.fill)
            target.number_format = source.number_format
            target.protection = copy(source.protection)
            target.alignment = copy(source.alignment)


def _write_sum_formulas(ws, last_data_row: int) -> None:
    sum_row = last_data_row + 1
    sum_columns = (8, 10, 11, 12, 13, 14, 15, 16, 17, 18, 19)
    for col in sum_columns:
        letter = get_column_letter(col)
        ws.cell(
            sum_row,
            col,
            f"=SUM({letter}{DATA_START_ROW}:{letter}{last_data_row})",
        )

    total_row = sum_row + 1
    ws.cell(total_row, 12, f"=SUM(L{sum_row}:S{sum_row})")


def export_count_sheet(rows: list[list], output_path: str | Path) -> None:
    if not rows:
        raise ValueError("Нет строк для листа подсчёта")

    template_path = resolve_count_template_path()
    workbook = openpyxl.load_workbook(template_path)
    worksheet = workbook.active

    for merged_range in list(worksheet.merged_cells.ranges):
        worksheet.unmerge_cells(str(merged_range))

    if worksheet.max_row > TEMPLATE_STYLE_ROW:
        worksheet.delete_rows(TEMPLATE_STYLE_ROW + 1, worksheet.max_row - TEMPLATE_STYLE_ROW)

    for offset, values in enumerate(rows):
        row_index = DATA_START_ROW + offset
        if row_index != TEMPLATE_STYLE_ROW:
            _copy_row_style(worksheet, TEMPLATE_STYLE_ROW, row_index, COUNT_DATA_COLUMNS)

        for col_index, value in enumerate(values[:COUNT_DATA_COLUMNS], start=1):
            worksheet.cell(row_index, col_index, value=value)

    last_data_row = DATA_START_ROW + len(rows) - 1
    _write_sum_formulas(worksheet, last_data_row)

    workbook.save(str(output_path))
